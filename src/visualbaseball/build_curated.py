"""Build the canonical game-sharded dataset from retained raw or legacy Excel."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import json
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from .collector import prepare_game
from .curated import file_sha256, validation_summary, write_game, write_schema
from .naver import NaverEnrichment


def _workbook_tables(path: Path) -> dict[str, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {}
        for sheet_name in ("Games", "Events", "Pitches"):
            if sheet_name not in workbook.sheetnames:
                tables[sheet_name] = []
                continue
            iterator = workbook[sheet_name].iter_rows(values_only=True)
            headers = [str(value or "") for value in next(iterator, ())]
            tables[sheet_name] = [{header: value for header, value in zip(headers, values)} for values in iterator]
        return tables
    finally:
        workbook.close()


def _baseline(rows: list[dict], season: int, source: Path, provenance: str) -> dict:
    pitch_ids = [str(row.get("pitch_id") or "") for row in rows]
    numeric = ("velocity_kmh", "px", "pz", "release_height_cm", "vertical_movement_cm", "horizontal_movement_cm")
    distributions = {}
    for field in numeric:
        values = sorted(float(row[field]) for row in rows if row.get(field) is not None)
        distributions[field] = {
            "n": len(values),
            "min": values[0] if values else None,
            "median": values[len(values) // 2] if values else None,
            "max": values[-1] if values else None,
        }
    return {
        "season": season, "source": str(source), "provenance": provenance,
        "source_sha256": file_sha256(source), "rows": len(rows),
        "games": len({str(row.get("game_id") or "") for row in rows}),
        "pitch_id_missing": sum(not value for value in pitch_ids),
        "pitch_id_duplicates": len(pitch_ids) - len(set(pitch_ids)),
        "columns": sorted({key for row in rows for key in row}), "distributions": distributions,
    }


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + f".{uuid4().hex}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def _build_raw(root: Path, storage_root: Path, season: int, game_id: str | None, force: bool) -> dict:
    raw_root = storage_root / "data" / "raw" / str(season)
    paths = sorted(raw_root.glob("*.json"))
    if game_id:
        paths = [path for path in paths if path.stem == game_id]
    if not paths:
        raise FileNotFoundError(f"no raw JSON found below {raw_root}")
    changed = pitches = 0
    for index, path in enumerate(paths, 1):
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        current_game_id = str(payload.get("gameData", {}).get("gameId") or path.stem)
        naver_path = storage_root / "data" / "raw" / "naver" / str(season) / f"{current_game_id}.json"
        enrichment = NaverEnrichment.from_dict(json.loads(naver_path.read_text(encoding="utf-8"))) if naver_path.exists() else None
        prepared = prepare_game(payload, season=season, naver_enrichment=enrichment)
        if not prepared.valid:
            raise ValueError(f"{current_game_id}: {prepared.message}")
        result = write_game(
            root, prepared.game, prepared.events, prepared.pitches, raw_payload=payload,
            provenance={"type": "visualbaseball_json", "path": path.relative_to(storage_root).as_posix()},
            force=force,
        )
        changed += int(result["changed"]); pitches += len(prepared.pitches)
        if index % 100 == 0 or index == len(paths):
            print(f"curated {index}/{len(paths)} games", flush=True)
    source_rows = sum((len((half_pa.get("pitches") or []))
                       for path in paths
                       for half in json.loads(path.read_text(encoding="utf-8-sig")).get("pbpData", [])
                       for half_pa in half.get("pas") or []), 0)
    baseline = {
        "season": season, "source": str(raw_root), "provenance": "visualbaseball_json",
        "files": len(paths), "rows": source_rows,
    }
    if game_id is None:
        _write_json(root / "data/curated/baselines" / f"season={season}.json", baseline)
    return {"games": len(paths), "changed": changed, "pitches": pitches, "baseline": baseline}


def _build_excel(root: Path, season: int, game_id: str | None, force: bool) -> dict:
    source = root / "exports" / f"visualbaseball_savant_{season}_latest.xlsx"
    if not source.is_file():
        raise FileNotFoundError(source)
    tables = _workbook_tables(source)
    pitches, events, games = tables["Pitches"], tables["Events"], tables["Games"]
    if game_id is None:
        _write_json(root / "data/curated/baselines" / f"season={season}.json", _baseline(pitches, season, source, "legacy_excel_one_time_conversion"))
    for row in pitches:
        if row.get("y0") is None:
            # Historical exports predate the provider's 2026-07 switch and x0/z0 are the 50 ft plane.
            row["y0"] = 50.0
    by_pitch, by_event = defaultdict(list), defaultdict(list)
    for row in pitches:
        by_pitch[str(row.get("game_id") or "")].append(row)
    for row in events:
        by_event[str(row.get("game_id") or "")].append(row)
    by_game = {str(row.get("game_id") or ""): row for row in games}
    ids = sorted(by_pitch)
    if game_id:
        ids = [value for value in ids if value == game_id]
    changed = 0
    for current_game_id in ids:
        game = by_game.get(current_game_id) or {
            "season": season, "game_date": f"{current_game_id[:4]}-{current_game_id[4:6]}-{current_game_id[6:8]}",
            "game_id": current_game_id, "validation_status": "MIGRATED_FROM_EXCEL",
        }
        game["season"] = season
        result = write_game(
            root, game, by_event[current_game_id], by_pitch[current_game_id],
            provenance={
                "type": "legacy_excel_one_time_conversion", "path": source.relative_to(root).as_posix(),
                "source_sha256": file_sha256(source), "raw_available": False,
                "coordinate_assumption": "missing y0 in historical export; x0/z0 documented and treated as y=50 ft",
            }, force=force,
        )
        changed += int(result["changed"])
    return {"games": len(ids), "changed": changed, "pitches": sum(len(by_pitch[value]) for value in ids)}


def build_curated(root: Path, season: int, *, storage_root: Path | None = None,
                  game_id: str | None = None, force: bool = False, validate: bool = False) -> dict:
    root, storage_root = root.resolve(), (storage_root or root).resolve()
    write_schema(root)
    raw_root = storage_root / "data" / "raw" / str(season)
    result = (_build_raw(root, storage_root, season, game_id, force)
              if any(raw_root.glob("*.json")) else _build_excel(root, season, game_id, force))
    if validate:
        summary = validation_summary(root, season)
        _write_json(root / "data/curated/validation" / f"season={season}.json", summary)
        if summary["pitch_id_duplicates"]:
            raise ValueError(f"season {season}: duplicate pitch_id={summary['pitch_id_duplicates']}")
        if summary["raw_pitches"] != summary["curated_pitches"] + summary["parse_excluded"]:
            raise ValueError(f"season {season}: raw/curated count reconciliation failed")
        for axis in ("plate_x_error", "plate_z_error"):
            distribution = summary[axis]
            if distribution["n"] and (
                distribution["median_abs_cm"] >= 1 or distribution["p95_abs_cm"] >= 2
            ):
                raise ValueError(f"season {season}: {axis} exceeded median <1 cm / p95 <2 cm")
        result["validation"] = summary
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path("."))
    parser.add_argument("--storage-root", type=Path)
    parser.add_argument("--season", type=int, required=True)
    parser.add_argument("--game-id")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--validate", action="store_true")
    args = parser.parse_args()
    print(json.dumps(build_curated(args.root, args.season, storage_root=args.storage_root,
                                   game_id=args.game_id, force=args.force, validate=args.validate), indent=2))


if __name__ == "__main__":
    main()
