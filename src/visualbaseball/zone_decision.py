"""League-execution ZA: (actual swing - expected swing) * (swing RV - take RV).

Chronological development and final holdouts evaluate event-decomposed values.
Three date-block cross-fits score every pitch without its own game's outcomes.
RE tables, models, event priors and support counts are fitted on training only.
"""
from __future__ import annotations
from collections import Counter, defaultdict
import argparse
import json
import hashlib
import platform
from importlib.metadata import version
from pathlib import Path
import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq
from sklearn.metrics import log_loss, brier_score_loss
from sklearn.isotonic import IsotonicRegression
from sklearn.model_selection import GroupKFold
from scipy.optimize import minimize, LinearConstraint
from openpyxl import load_workbook
from . import plate_decision_v1 as old
from .pitch_arsenal import _load_batter_hands, _resolved_batter_stance
from .swing_take import _eligible, _relative_location, _state
from .zone_awareness_v2 import _team_history

REGIONS = ('heart', 'shadow_in', 'shadow_out', 'chase', 'waste')
EVENTS = ('Whiff', 'Foul', 'InPlay', 'Ball', 'CalledStrike', 'HBP')
NUMERIC = old.BASE_NUMERIC + old.MOVEMENT_NUMERIC
MODEL_VERSION = 'za5-state-transitions'
CONTRACT = {
 'za_raw': '100 * mean((S - p_swing) * (V_swing - V_take)); runs per 100 pitches',
 'raw_dv': 'sum((S - p_swing) * (V_swing - V_take)); cumulative runs',
 'dv_per_100': 'same as za_raw; retained for backward-compatible consumers',
 'zone_judgment_raw': '100 * mean((S - p_swing) * (2*p_called_strike_if_take - 1)); percentage points',
 'swing_aggression': '100 * mean(S - p_swing); percentage points, tendency only',
 'za_percentile': 'midrank percentile among season hitters with at least 300 eligible pitches',
 'region_contributions': '100 * sum(D in region/action) / ALL eligible player pitches; additive',
}
LIMITATIONS = [
 '리그 평균 실행 능력을 기준으로 추정한 의사결정 가치이며 개인별 최적 판단의 정답이 아닙니다.',
 '관측하지 못한 반대 행동과 누락된 투구 특성에 따른 선택 편향이 남습니다. 실제 타구속도·발사각·해당 투구의 안타/홈런은 판단 점수의 입력이 아닙니다.',
 '시즌 표시값은 날짜 블록 교차적합으로 해당 경기 결과를 제외하지만 다른 블록의 미래 경기를 사용할 수 있습니다. 순수한 사전 예측 성능은 별도 시간 분리 평가에서 확인합니다.',
 '표본 부족 구간은 상위 조건의 결과 분포로 완화합니다. 반대 선택의 정확성과 누락된 실행 능력·번트 의도의 영향까지 검증된 것은 아닙니다.',
 '95% 구간은 경기 단위 재표집이며 학습 모델을 고정합니다. 모델 추정 오차까지 포함하는 전체 신뢰구간이 아닙니다.',
]


def file_hash(path):
 with path.open('rb') as stream:
  return hashlib.file_digest(stream, 'sha256').hexdigest()


def reliable_halves(rows, events):
 """Use timed, trusted non-pitch runs; never assign untimed score repairs to a pitch."""
 halves, timeline = defaultdict(list), defaultdict(list)
 bad = defaultdict(set)
 for r in rows:
  halves[(r['game_id'],r['inning'],r['inning_half'])].append(r)
 for e in events:
  key=(e['game_id'],e['inning'],e['inning_half'])
  runs=float(e.get('runs_on_event') or 0)
  score_change=sum(float(e.get(f'{team}_score_after') or 0)-float(e.get(f'{team}_score_before') or 0) for team in ('away','home'))
  if e.get('event_code') in ('OFFICIAL_LINESCORE_RECONCILIATION','SOURCE_SCORE_SNAPSHOT') and score_change:
   bad[key].add('untimed_score_repair')
  if e.get('event_type')!='pitch' and (runs or score_change):
   if e.get('parse_status')=='ok' and runs==score_change and runs>=0:
    timeline[key].append((e['event_seq'],runs,None))
   else:
    bad[key].add('unresolved_nonpitch_runs')
 accepted=[];reasons=Counter();complete=0;known_runs=0.
 for key,items in halves.items():
  items.sort(key=lambda r:r['event_seq'])
  # Score snapshots must agree with the represented pitch/non-pitch runs.
  start,end=items[0],items[-1]
  expected=sum(float(end.get(f'{team}_score_after') or 0)-float(start.get(f'{team}_score_before') or 0) for team in ('away','home'))
  represented=sum(float(r.get('runs_on_pitch') or 0) for r in items)
  represented+=sum(v for seq,v,_ in timeline[key] if start['event_seq']<=seq<=end['event_seq'])
  if all(f'{t}_score_{w}' in r for r,w in ((start,'before'),(end,'after')) for t in ('away','home')) and expected!=represented:
   bad[key].add('score_timeline_mismatch')
  if bad[key]:
   reasons.update(bad[key]);continue
  is_complete=end.get('outs_after')==3
  complete+=int(is_complete)
  remaining=0.
  ordered=timeline[key]+[(r['event_seq'],float(r.get('runs_on_pitch') or 0),r) for r in items]
  for _,runs,r in sorted(ordered,key=lambda item:item[0],reverse=True):
   remaining+=runs
   if r is not None:
    r['_runs_to_end']=remaining;r['_re_complete']=is_complete
  known_runs+=sum(v for _,v,_ in timeline[key])
  accepted.extend(items)
 return accepted,{'halves':len(halves),'excluded_halves':sum(bool(bad[k]) for k in halves),
   'excluded_pitches':len(rows)-len(accepted),'re_complete_halves':complete,
   'reason_counts':dict(reasons),'included_timed_nonpitch_runs':known_runs,
   'policy':'Unresolved scoring halves are excluded; incomplete innings do not train RE. No invented scoring timestamps.'}


def decision_value(swing, probability, swing_value, take_value):
 return (np.asarray(swing) - np.asarray(probability)) * (np.asarray(swing_value) - np.asarray(take_value))


def region(row):
 d = max(abs(row['x_relative']), abs(row['z_relative']))
 return 'heart' if d <= 2/3 else 'shadow_in' if d <= 1 else 'shadow_out' if d <= 4/3 else 'chase' if d <= 2 else 'waste'


def outcome(row):
 # Classify by this pitch's call, never a nonterminal pitch's eventual PA result.
 code = str(row.get('pitch_call_code') or '').upper()
 if row.get('is_pa_terminal') and (str(row.get('pa_type') or '').lower() == 'hbp' or row.get('pa_result') == '사구'):
  return 'HBP'
 return {'S':'Whiff', 'F':'Foul', 'X':'InPlay', 'B':'Ball', 'T':'CalledStrike'}.get(code)


def load_rows(root, season, storage_root=None):
 cache = root / '.cache' / f'za_source_{season}.parquet'
 cache.parent.mkdir(parents=True,exist_ok=True)
 source = (storage_root or root) / 'data/processed/pitches.parquet'
 hashes={}
 if season == 2026:
  rows = pq.read_table(source).to_pylist()
  event_source=source.with_name('events.parquet')
  events=pq.read_table(event_source).to_pylist()
  hashes={p.name:file_hash(p) for p in (source,event_source)}
 else:
  source=root/'exports'/f'visualbaseball_savant_{season}_latest.xlsx'
  fingerprint=file_hash(source);hashes[source.name]=fingerprint
  marker=cache.with_suffix('.sha256');event_cache=cache.with_name(f'za_events_{season}.parquet')
  if cache.exists() and event_cache.exists() and marker.exists() and marker.read_text()==fingerprint:
   rows=pq.read_table(cache).to_pylist();events=pq.read_table(event_cache).to_pylist()
  else:
   workbook=load_workbook(source,read_only=True,data_only=True)
   try:
    tables=[]
    for sheet in ('Pitches','Events'):
     iterator=workbook[sheet].iter_rows(values_only=True);headers=next(iterator)
     tables.append([dict(zip(headers,values)) for values in iterator])
    rows,events=tables
   finally:
    workbook.close()
   rows=[r for r in rows if r.get('season')==season]
   pq.write_table(pa.Table.from_pylist(rows),cache)
   pq.write_table(pa.Table.from_pylist(events),event_cache)
   marker.write_text(fingerprint,encoding='utf-8')
 rows=[r for r in rows if int(r.get('season') or season)==season]
 rows,quality=reliable_halves(rows,events)
 hands = _load_batter_hands(root, season)
 valid, excluded = [], Counter()
 for r in rows:
  if not _eligible(r) or not r.get('batter_id') or not r.get('batter_name') or outcome(r) is None:
   excluded['invalid_state_location_action_or_identity'] += 1
   continue
  r['x_relative'], r['z_relative'] = _relative_location(r)
  r['decision_type'] = 'Swing' if outcome(r) in EVENTS[:3] else 'Take'
  r['batter_stance'] = _resolved_batter_stance(r, hands)
  r['event'] = outcome(r)
  r['region'] = region(r)
  valid.append(r)
 movement = old._movement_adjust(valid, root, season)
 # Retain pre-pitch features, transitions, identity and training target only.
 keep = set(NUMERIC + old.CATEGORICAL + ('game_id','game_date','season','batter_id','batter_name','batter_team','inning_half','event','region','decision_type','_runs_to_end','_re_complete','runs_on_pitch'))
 keep.update(f'{k}_{w}' for k in ('base_state_code','outs','balls','strikes') for w in ('before','after'))
 valid = [{k:v for k,v in r.items() if k in keep} for r in valid]
 for p in (root/'data/batter_handedness.json',root/'data/park_adjustments'/f'{season}_VB_Park_Adjustment_v1.0.xlsx'):
  if p.exists():hashes[p.name]=file_hash(p)
 return sorted(valid, key=lambda r:r['game_id']), {'source':source.name, 'sha256':hashes,'quality':quality,'excluded':dict(excluded), 'movement':movement, 'unknown_stance':sum(not r['batter_stance'] for r in valid),'latest_game':max(r['game_id'] for r in valid)}


def walk_state(s):
 bases,outs,_,_=s
 # Only forced runners advance; a bases-loaded walk scores exactly one run.
 after=bases|1
 if bases&1:after|=2
 if bases&3==3:after|=4
 return (after,outs,0,0),int(bases==7)


class RunExpectancy:
 def __init__(self, rows):
  full, base, outs = defaultdict(list), defaultdict(list), defaultdict(list)
  for r in rows:
   if not r.get('_re_complete',True):continue
   s = _state(r, 'before'); v = r['_runs_to_end']
   full[s].append(v); base[s[:2]].append(v); outs[s[1]].append(v)
  self.outs = {k:float(np.mean(v)) for k,v in outs.items()}
  self.base = {k:(sum(v)+50*self.outs[k[1]])/(len(v)+50) for k,v in base.items()}
  self.full = {k:(sum(v)+50*self.base[k[:2]])/(len(v)+50) for k,v in full.items()}
  if not full:raise ValueError('No complete, reliable innings for run expectancy')
  # Weighted least-squares projection of the entire RE table, not clipping RVs.
  # Ball, strike, strikeout and forced-walk transitions constrain one shared table.
  states=[(b,o,ball,strike) for b in range(8) for o in range(3) for ball in range(4) for strike in range(3)]
  index={s:i for i,s in enumerate(states)}
  prior=np.array([self.value(s) for s in states]);weights=np.array([len(full[s])+50 for s in states],dtype=float)
  weights/=weights.max();constraints=[];bounds=[]
  def edge(a,b,limit=0):
   row=np.zeros(len(states));row[index[a]]=1
   if b[1]<3:row[index[b]]-=1
   constraints.append(row);bounds.append(limit)
  for s in states:
   b,o,ball,strike=s
   if ball<3:edge(s,(b,o,ball+1,strike))
   if strike<2:edge((b,o,ball,strike+1),s)
   if o<2:edge((b,o+1,0,0),s)
   after,runs=walk_state(s);edge(s,after,runs)
  matrix=np.array(constraints);upper=np.array(bounds)
  fitted=minimize(lambda x:float(np.sum(weights*(x-prior)**2)/2),prior,
   jac=lambda x:weights*(x-prior),method='SLSQP',bounds=[(0,None)]*len(states),
   constraints=[LinearConstraint(matrix,-np.inf,upper)],options={'ftol':1e-10,'maxiter':500})
  if not fitted.success or np.max(matrix@fitted.x-upper)>1e-7:
   raise RuntimeError(f'Constrained RE failed: {fitted.message}')
  self.full=dict(zip(states,fitted.x))
  self.diagnostics={'states':len(states),'constraints':len(bounds),'max_constraint_violation':float(max(0,np.max(matrix@fitted.x-upper))),
   'weighted_adjustment_rmse':float(np.sqrt(np.average((fitted.x-prior)**2,weights=weights)))}
 def value(self, s):
  if s[1] >= 3: return 0.
  return self.full.get(s, self.base.get(s[:2], self.outs.get(s[1], 0.)))
 def target(self, rows):
  return np.array([float(r.get('runs_on_pitch') or 0)+self.value(_state(r,'after'))-self.value(_state(r,'before')) for r in rows])

 def event_value(self,s,event):
  b,o,balls,strikes=s;runs=0
  if event in ('Ball','HBP') and (balls==3 or event=='HBP'):after,runs=walk_state(s)
  elif event=='Ball':after=(b,o,balls+1,strikes)
  elif event=='Foul' and strikes==2:return 0.
  elif event in ('Whiff','CalledStrike') and strikes==2:after=(b,o+1,0,0)
  elif event in ('Whiff','CalledStrike','Foul'):after=(b,o,balls,strikes+1)
  else:raise ValueError(f'Unsupported deterministic event: {event}')
  return float(runs+self.value(after)-self.value(s))


def encode(train, test):
 # Unknown categories are missing, never mapped to another known category.
 cols_a, cols_b = [], []
 for f in NUMERIC:
  cols_a.append([old._safe_float(r.get(f)) for r in train]); cols_b.append([old._safe_float(r.get(f)) for r in test])
 for f in old.CATEGORICAL:
  mapping = {v:i for i,v in enumerate(sorted({str(r.get(f) or '') for r in train}))}
  cols_a.append([mapping[str(r.get(f) or '')] for r in train]); cols_b.append([mapping.get(str(r.get(f) or ''),np.nan) for r in test])
 return np.column_stack(cols_a), np.column_stack(cols_b)


def support_key(r):
 return (round(r['x_relative']*2),round(r['z_relative']*2),r['balls_before'],r['strikes_before'])


def detailed_support_key(r):
 def band(key,width):
  value=old._safe_float(r.get(key))
  return int(np.floor(value/width)) if np.isfinite(value) else None
 return (*support_key(r),r.get('pitch_type'),r.get('batter_stance'),
         band('velocity_kmh',10),band('adjusted_hb_cm',10),band('adjusted_ivb_cm',10))


def probability_prior(train,test,actions,events,action,indices):
 """Training-only hierarchy: count/region/type/stance -> count/region -> action."""
 def coarse(r):return (r['balls_before'],r['strikes_before'],r['region'])
 def fine(r):return (*coarse(r),r.get('pitch_type'),r.get('batter_stance'))
 global_counts=np.ones(len(indices));parent=defaultdict(lambda:np.zeros(len(indices)));local=defaultdict(lambda:np.zeros(len(indices)))
 for i in np.flatnonzero(actions==action):
  j=list(indices).index(events[i]);global_counts[j]+=1;parent[coarse(train[i])][j]+=1;local[fine(train[i])][j]+=1
 global_p=global_counts/global_counts.sum()
 result=[]
 for r in test:
  counts=parent[coarse(r)];p=(counts+50*global_p)/(counts.sum()+50)
  counts=local[fine(r)];result.append((counts+50*p)/(counts.sum()+50))
 return np.array(result)


def fit_model(model, features, target):
 # All-missing training columns carry no information. A constant lets histogram
 # binning ignore them, including in action-specific and calibration subsets.
 missing=np.isnan(features).all(axis=0)
 if missing.any():
  features=features.copy();features[:,missing]=0
 return model.fit(features,target)


def fit_predict(train, test, candidate=True, calibration=True, support_prior=50):
 a,b = encode(train,test)
 re = RunExpectancy(train); target = re.target(train)
 actions = np.array([r['decision_type']=='Swing' for r in train],dtype=int)
 events = np.array([EVENTS.index(r['event']) for r in train])
 propensity = fit_model(old._classifier(len(NUMERIC)),a,actions)
 raw_p = propensity.predict_proba(b)[:,list(propensity.classes_).index(1)]
 p=raw_p.copy()
 calibration_applied=False
 if calibration:
  groups=np.array([r['game_id'] for r in train]);folds=min(3,len(set(groups)))
  if folds<2:raise ValueError('Calibration requires at least two training games')
  oof=np.empty(len(train))
  for fit,held in GroupKFold(folds).split(a,actions,groups):
   model=fit_model(old._classifier(len(NUMERIC)),a[fit],actions[fit])
   oof[held]=model.predict_proba(a[held])[:,list(model.classes_).index(1)]
  days=sorted({r['game_id'][:8] for r in train});cut=days[max(1,int(len(days)*.8))-1]
  fit_mask=np.array([r['game_id'][:8]<=cut for r in train]);held_mask=~fit_mask
  if held_mask.any() and len(set(actions[fit_mask]))==2:
   calibrator=IsotonicRegression(out_of_bounds='clip',y_min=1e-6,y_max=1-1e-6).fit(oof[fit_mask],actions[fit_mask])
   corrected=calibrator.predict(oof[held_mask]);truth=actions[held_mask]
   calibration_applied=(log_loss(truth,corrected,labels=[0,1])<log_loss(truth,oof[held_mask],labels=[0,1]) and brier_score_loss(truth,corrected)<=brier_score_loss(truth,oof[held_mask]))
   if calibration_applied:
    calibrator.fit(oof,actions);p=calibrator.predict(raw_p)
 probs = np.zeros((len(test),6)); direct=np.zeros((len(test),2)); staged=direct.copy();unsmoothed=direct.copy()
 counts = Counter((support_key(r),r['decision_type']) for r in train)
 support = np.array([[counts[(support_key(r),action)] for action in ('Take','Swing')] for r in test])
 fine_counts=Counter((detailed_support_key(r),r['decision_type']) for r in train)
 detailed=np.array([[fine_counts[(detailed_support_key(r),action)] for action in ('Take','Swing')] for r in test])
 for act, indices in ((0,range(3,6)),(1,range(3))):
  mask = actions==act
  direct[:,act] = fit_model(old._regressor(len(NUMERIC)),a[mask],target[mask]).predict(b)
  clf = fit_model(old._classifier(len(NUMERIC)),a[mask],events[mask])
  pred = clf.predict_proba(b)
  for i,label in enumerate(clf.classes_): probs[:,int(label)]=pred[:,i]
  raw_probs=probs[:,list(indices)].copy()
  if act==1 and support_prior:
   weight=detailed[:,act]/(detailed[:,act]+support_prior)
   prior=probability_prior(train,test,actions,events,act,indices)
   probs[:,list(indices)]=weight[:,None]*raw_probs+(1-weight[:,None])*prior
  if not candidate: continue
  # Event RV priors back off event x count x base/out -> event x count -> event.
  for ev in indices:
   em = events==ev
   if ev!=2:
    values=np.array([re.event_value(_state(r,'before'),EVENTS[ev]) for r in test])
    staged[:,act]+=probs[:,ev]*values
    unsmoothed[:,act]+=raw_probs[:,list(indices).index(ev)]*values
    continue
   global_mean = float(np.mean(target[em])) if em.any() else float(np.mean(target[mask]))
   by_count, by_state = defaultdict(list), defaultdict(list)
   for j in np.flatnonzero(em):
    r=train[j]; by_count[(r['balls_before'],r['strikes_before'])].append(target[j]); by_state[_state(r,'before')].append(target[j])
   cm={k:(sum(v)+50*global_mean)/(len(v)+50) for k,v in by_count.items()}
   sm={k:(sum(v)+50*cm[k[2:]])/(len(v)+50) for k,v in by_state.items()}
   prior=np.array([sm.get(_state(r,'before'),cm.get(_state(r,'before')[2:],global_mean)) for r in test])
   values=prior
   if ev==2 and em.sum()>=160:
    model=fit_model(old._regressor(len(NUMERIC)),a[em],target[em])
    local=Counter(support_key(train[j]) for j in np.flatnonzero(em))
    n=np.array([local[support_key(r)] for r in test]); weight=n/(n+50.)
    values=weight*model.predict(b)+(1-weight)*prior
   staged[:,act]+=probs[:,ev]*values
   unsmoothed[:,act]+=raw_probs[:,list(indices).index(ev)]*values
 return {'p':p,'raw_p':raw_p,'direct':direct,'staged':staged,'unsmoothed':unsmoothed,'probs':probs,'support':support,
         'detailed_support':detailed,'target':re.target(test),'re_diagnostics':re.diagnostics,'calibration_applied':calibration_applied}


def evaluation_metrics(test,pred):
 action=np.array([r['decision_type']=='Swing' for r in test],dtype=int); idx=np.arange(len(test))
 metrics={}
 for name in ('direct','staged','unsmoothed'):
  errors=(pred[name][idx,action]-pred['target'])**2
  metrics[name]={'mse':float(errors.mean()),'swing_mse':float(errors[action==1].mean()),'take_mse':float(errors[action==0].mean())}
 event_metrics={}
 labels=np.array([EVENTS.index(r['event']) for r in test])
 for name,act,inds in (('swing',1,list(range(3))),('take',0,list(range(3,6)))):
  mask=action==act
  event_metrics[name+'_log_loss']=float(log_loss(labels[mask],pred['probs'][mask][:,inds],labels=inds))
 calibration=[]
 for low in np.arange(0,1,.1):
  mask=(pred['p']>=low)&(pred['p']<low+.1)
  if mask.any():calibration.append({'low':r6(low),'n':int(mask.sum()),'predicted':float(pred['p'][mask].mean()),'observed':float(action[mask].mean())})
 return {'pitches':len(test),'swing_log_loss':float(log_loss(action,pred['p'])),'swing_brier':float(brier_score_loss(action,pred['p'])),
  'raw_swing_log_loss':float(log_loss(action,pred['raw_p'])),'raw_swing_brier':float(brier_score_loss(action,pred['raw_p'])),
  'action_value':metrics,'event_probability':event_metrics,'calibration_bins':calibration,'re':pred['re_diagnostics'],'calibration_applied':pred['calibration_applied']}


def temporal_evaluation(rows):
 dates=sorted({r['game_id'][:8] for r in rows})
 if len(dates)<10:raise ValueError('At least ten dates are required for temporal validation')
 dev_date,test_date=dates[int(len(dates)*.6)],dates[int(len(dates)*.8)]
 train=[r for r in rows if r['game_id'][:8]<dev_date]
 dev=[r for r in rows if dev_date<=r['game_id'][:8]<test_date]
 test=[r for r in rows if r['game_id'][:8]>=test_date]
 print('  Development',dev_date,len(train),len(dev),flush=True)
 predicted=fit_predict(train,dev)
 development=evaluation_metrics(dev,predicted)
 # Same algorithm as season scoring. Optional probability calibration is gated
 # inside each training set, never with a scored/development/final-test outcome.
 settings={'calibration':True,'support_prior':50}
 print('  Final untouched test',test_date,len(test),settings,flush=True)
 predicted=fit_predict(train+dev,test,**settings)
 final=evaluation_metrics(test,predicted)
 return {'selected':'staged','settings':settings,'development_start':dev_date,'split_date':test_date,
  'train_pitches':len(train)+len(dev),'test_pitches':len(test),'development':development,'final_test':final,
  **{k:final[k] for k in ('swing_log_loss','swing_brier','action_value','event_probability')},
  'selection_rule':'60/20/20 dates: development diagnostics, then final evaluation with frozen algorithm. Each fit selects optional calibration using only training-game OOF probabilities. Final 20% never selects settings. Direct and unpooled RVs are diagnostic baselines.',
  'counterfactual_validation':'Observed-action errors and conditional intervals do not identify unobserved opposite-action outcomes.'}


def r6(x): return round(float(x),6)
def mean(items,key,scale=1): return r6(scale*np.mean([r[key] for r in items])) if items else None


def profile_summary(items):
 n=len(items); first=items[0]
 total=sum(r['dv'] for r in items)
 s={'season':first['season'],'batter_id':str(first['batter_id']),'batter_name':first['batter_name'],'team':_team_history(items),'pitches_seen':n,'qualified_300':n>=300,'za_raw':r6(100*total/n),'dv_per_100':r6(100*total/n),'raw_dv':r6(total),'swing_aggression':r6(100*np.mean([r['swing']-r['p_swing'] for r in items])),'zone_judgment_raw':mean(items,'judgment',100),'za_percentile':None,'low_opposite_support_pitches':sum(r['opposite_support']<30 for r in items)}
 s['low_opposite_support_pct']=r6(100*s['low_opposite_support_pitches']/n)
 s['za_ci_low']=s['za_ci_high']=None
 games=defaultdict(list)
 for r in items:games[r['game_id']].append(r['dv'])
 s['games_seen']=len(games)
 if n>=300 and len(games)>1:
  sums=np.array([sum(v) for v in games.values()]);counts=np.array([len(v) for v in games.values()])
  sampled=np.random.default_rng(old.RANDOM_STATE).integers(0,len(games),(1000,len(games)))
  estimates=100*sums[sampled].sum(axis=1)/counts[sampled].sum(axis=1)
  s['za_ci_low'],s['za_ci_high']=map(r6,np.quantile(estimates,[.025,.975]))
 for action in ('swing','take'):
  selected=[r for r in items if r['swing']==(action=='swing')]
  s[action+'_pitches']=len(selected)
  s[action+'_decision_value_per_100']=r6(100*sum(r['dv'] for r in selected)/n)
 for reg in REGIONS:
  selected=[r for r in items if r['region']==reg]
  s[reg+'_pitches']=len(selected); s[reg+'_raw_dv']=r6(sum(r['dv'] for r in selected))
  s[reg+'_decision_value_per_100']=r6(100*sum(r['dv'] for r in selected)/n)
  for action in ('swing','take'):
   value=sum(r['dv'] for r in selected if r['swing']==(action=='swing'))
   s[f'{reg}_{action}_decision_value_per_100']=r6(100*value/n)
 return s


def cell_summary(items):
 return {'n':len(items),'low_opposite_support_pct':r6(100*np.mean([r['opposite_support']<30 for r in items])),'raw_dv':r6(sum(r['dv'] for r in items)),'dv100':mean(items,'dv',100),'za_raw':mean(items,'dv',100),'delta':mean(items,'delta_v'),'swing_pct':mean(items,'swing',100),'expected_swing_pct':mean(items,'p_swing',100),'p_zone_pct':mean(items,'p_zone',100),'zone_judgment_pct':r6(100*np.mean([r['p_zone'] if r['swing'] else 1-r['p_zone'] for r in items])),'expected_zone_judgment_pct':r6(100*np.mean([r['p_swing']*r['p_zone']+(1-r['p_swing'])*(1-r['p_zone']) for r in items])),'expected_swing_rv':mean(items,'v_swing'),'expected_take_rv':mean(items,'v_take'),**{f'p_{e}':mean(items,f'p_{e}',100) for e in EVENTS}}


def write_web(root,season,pitches,report):
 dest=root/'web/data/zone_awareness'/str(season); dest.mkdir(parents=True,exist_ok=True)
 by_batter=defaultdict(list)
 for r in pitches: by_batter[str(r['batter_id'])].append(r)
 players=[profile_summary(items) for items in by_batter.values()]
 scores=np.array([p['za_raw'] for p in players if p['qualified_300']])
 for p in players: p['za_percentile']=r6(100*(np.sum(scores<p['za_raw'])+.5*np.sum(scores==p['za_raw']))/len(scores)) if len(scores) else None
 def dump(path,payload): path.write_text(json.dumps(payload,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\n',encoding='utf-8')
 dump(dest/'leaderboard.json',{'schema_version':5,'model_version':MODEL_VERSION,'season':season,'minimum_pitches':300,'qualified_batters':len(scores),'players':players,'metric_contract':CONTRACT,'selected_value_model':report['validation']['selected'],'data_quality':report['source']['quality'],'settings':report['validation']['settings']})
 dump(dest/'teams.json',{'season':season,'teams':{p['batter_id']:p['team'] for p in players}})
 shards=defaultdict(dict)
 for p in players:
  bid=p['batter_id']; items=by_batter[bid]; cells=defaultdict(list)
  for r in items:
   x,z=r['x_relative'],r['z_relative']
   if abs(x)<=2.5 and abs(z)<=2.5: cells[(round(x/.5)*.5,round(z/.5)*.5)].append(r)
  shards[bid[:2] if bid[0].isdigit() else 'other'][bid]={'summary':p,'overall':cell_summary(items),'grid':[{'x':x,'z':z,**cell_summary(rs)} for (x,z),rs in sorted(cells.items())]}
 pd=dest/'players';pd.mkdir(exist_ok=True)
 for path in pd.glob('*.json'): path.unlink()
 for shard,entries in shards.items(): dump(pd/f'{shard}.json',{'schema_version':5,'season':season,'players':entries})
 catalog=root/'web/data/zone_awareness/index.json'
 previous=json.loads(catalog.read_text(encoding='utf-8')) if catalog.exists() else {'seasons':[]}
 previous.update({'schema_version':5,'seasons':sorted(set(previous['seasons'])|{season},reverse=True)})
 previous['default_season']=max(previous['seasons']);dump(catalog,previous)
 old._write_csv(root/'exports'/f'zone_decision_players_{season}.csv',players)
 return players


def build_zone_decision(root,season=2026,storage_root=None):
 print('ZA season',season,flush=True)
 rows,source=load_rows(root,season,storage_root)
 validation=temporal_evaluation(rows)
 selected=validation['selected'];print('  Selected:',selected,flush=True)
 dates=np.array(sorted({r['game_id'][:8] for r in rows})); result=[]; fold_meta=[]
 for fold,block in enumerate(np.array_split(dates,3)):
  held=set(block);train=[r for r in rows if r['game_id'][:8] not in held];test=[r for r in rows if r['game_id'][:8] in held]
  print('  Scoring block',fold+1,len(test),flush=True)
  # Fix score-model settings before seeing any season outcomes. The separate
  # development gate is for prospective evaluation, never a feedback path from
  # a scored game's outcome into that game's hyperparameters.
  pred=fit_predict(train,test,calibration=True,support_prior=50)
  action=np.array([r['decision_type']=='Swing' for r in test],dtype=int)
  values=pred[selected]; dv=decision_value(action,pred['p'],values[:,1],values[:,0])
  for i,r in enumerate(test):
   pzone=pred['probs'][i,4]
   r.update({'swing':int(action[i]),'p_swing':float(pred['p'][i]),'raw_p_swing':float(pred['raw_p'][i]),'p_zone':float(pzone),'judgment':float((action[i]-pred['p'][i])*(2*pzone-1)),'v_swing':float(values[i,1]),'v_take':float(values[i,0]),'delta_v':float(values[i,1]-values[i,0]),'dv':float(dv[i]),'opposite_support':int(pred['detailed_support'][i,1-action[i]]),'coarse_opposite_support':int(pred['support'][i,1-action[i]]),'fold':fold})
   for j,e in enumerate(EVENTS): r[f'p_{e}']=float(pred['probs'][i,j])
  result.extend(test);fold_meta.append({'start':str(block[0]),'end':str(block[-1]),'pitches':len(test),'re':pred['re_diagnostics'],'calibration_applied':pred['calibration_applied']})
 # One report: observed fit, period reproducibility and opposite-action support.
 periods=defaultdict(lambda:defaultdict(list))
 for r in result: periods[r['fold']][str(r['batter_id'])].append(r['dv'])
 stability=[]
 for a,b in ((0,1),(1,2)):
  ids=[k for k,v in periods[a].items() if len(v)>=150 and len(periods[b].get(k,[]))>=150]
  x=[np.mean(periods[a][k]) for k in ids];y=[np.mean(periods[b][k]) for k in ids]
  corr=float(np.corrcoef(x,y)[0,1]) if len(ids)>2 and np.std(x)>0 and np.std(y)>0 else None
  stability.append({'blocks':[a+1,b+1],'batters_150_pitches_each':len(ids),'pearson_r':corr,'interpretation':'descriptive repeatability; cross-fit training overlaps, not independent prospective validation'})
 support={reg:{'pitches':sum(r['region']==reg for r in result),'opposite_action_under_30':sum(r['region']==reg and r['opposite_support']<30 for r in result)} for reg in REGIONS}
 report={'schema_version':5,'model_version':MODEL_VERSION,'season':season,'source':source,'pitches':len(result),'validation':validation,'period_reproducibility':stability,'opposite_action_support':support,
  'support_definition':'Opposite-action counts: normalized 0.5 location cell x count x pitch type x stance x 10 km/h velocity x 10 cm HB/IVB bins. Diagnostic neighborhood, not proof of causal overlap.',
  'score_settings':{'calibration':True,'support_prior':50,'policy':'Fixed before outcome inspection; game-block exclusions apply to calibration, RE, models and priors.'},
  'crossfit_blocks':fold_meta,'metric_contract':CONTRACT,'shrinkage':'Constrained shared RE; exact ordinary event transitions; sparse Swing probabilities blend to training-only count/region/type/stance priors with n/(n+50). InPlay RV retains state shrinkage.',
  'reproducibility':{'python':platform.python_version(),'packages':{p:version(p) for p in ('numpy','pandas','pyarrow','scikit-learn','scipy','openpyxl')},'source_sha256':{p.name:file_hash(p) for p in (Path(__file__),Path(old.__file__),Path(__file__).with_name('swing_take.py'),Path(__file__).with_name('pitch_arsenal.py'))}},
  'limitations':LIMITATIONS}
 players=write_web(root,season,result,report)
 report['batters']=len(players)
 path=root/'data/processed'/f'zone_decision_report_{season}.json';path.parent.mkdir(parents=True,exist_ok=True)
 path.write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
 # Compact pitch evidence is reproducible; not committed as a large binary.
 pq.write_table(pa.Table.from_pylist([{k:v for k,v in r.items() if k not in ('adjusted_hb_cm','adjusted_ivb_cm')} for r in result]),root/'.cache'/f'zone_decision_pitches_{season}.parquet')
 return report

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--root',default='.');parser.add_argument('--seasons',nargs='+',type=int,default=[2024,2025,2026]);args=parser.parse_args()
 for year in args.seasons: build_zone_decision(Path(args.root).resolve(),year)
