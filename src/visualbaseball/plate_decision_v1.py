"""Shared plate-decision model components used by zone_decision.

The p_swing / p_zone classifier and RV regressor settings, feature encoders and the
park-adjusted movement path (``_movement_adjust``, read only when the season's
park-adjustment workbook exists). The cross-fitted v1 research pipeline that used to
live here was superseded by zone_decision for every published season and removed.
"""
from __future__ import annotations

import csv
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.ensemble import HistGradientBoostingClassifier, HistGradientBoostingRegressor

from .curated import _number
from .pitch_types import pitch_code


RANDOM_STATE = 20260903
PARK_FACTOR_CODES = ("FF", "SI", "FC", "SL", "CH", "CU", "FS")
PARK_FACTOR_CODE = {"FT": "SI", "ST": "SL"}

BASE_NUMERIC = (
    "x_relative", "z_relative", "balls_before", "strikes_before", "outs_before",
    "base_state_code_before", "velocity_kmh", "release_height_cm",
)
MOVEMENT_NUMERIC = ("adjusted_hb_cm", "adjusted_ivb_cm")
CATEGORICAL = ("pitch_type", "batter_stance", "stadium")
PZONE_NUMERIC = ("x_relative", "z_relative", "sz_top", "sz_bottom")


def _classifier(categorical_start: int | None = None) -> HistGradientBoostingClassifier:
    return HistGradientBoostingClassifier(
        learning_rate=0.07, max_iter=130, max_leaf_nodes=20,
        min_samples_leaf=80, l2_regularization=1.5, random_state=RANDOM_STATE,
        categorical_features=(list(range(categorical_start, categorical_start + len(CATEGORICAL)))
                              if categorical_start is not None else None),
    )


def _regressor(categorical_start: int) -> HistGradientBoostingRegressor:
    return HistGradientBoostingRegressor(
        learning_rate=0.06, max_iter=130, max_leaf_nodes=18,
        min_samples_leaf=80, l2_regularization=2.0, random_state=RANDOM_STATE,
        categorical_features=list(range(categorical_start, categorical_start + len(CATEGORICAL))),
    )


def _safe_float(value) -> float:
    try:
        value = float(value)
        return value if np.isfinite(value) else np.nan
    except (TypeError, ValueError):
        return np.nan


def _encode_numeric(rows: list[dict], fields: tuple[str, ...]) -> np.ndarray:
    return np.column_stack([
        np.array([_safe_float(row.get(field)) for row in rows], dtype=float) for field in fields
    ])


def _stadium(value) -> str:
    name = str(value or "").replace(" ", "")
    aliases = (
        ("고척", "고척"), ("광주", "광주"), ("대구", "대구"),
        ("대전", "대전"), ("한밭", "대전"), ("문학", "문학"), ("인천", "문학"),
        ("사직", "사직"), ("수원", "수원"), ("잠실", "잠실"), ("창원", "창원"),
    )
    return next((canonical for token, canonical in aliases if token in name), name)


def _load_park_factors(root: Path, season: int) -> dict[tuple[str, str], tuple[float, float]]:
    """Return {(stadium, pitch code): (HB offset cm, IVB offset cm)}.

    The 2022-25 workbooks use a fixed seven-column pitch order.  Some supplied
    header cells contain duplicate labels, so positions are used deliberately;
    the intact 2023 and 2025 files establish the shared order.
    """
    source = root / "data" / "park_adjustments" / f"{season}_VB_Park_Adjustment_v1.0.xlsx"
    if not source.exists():
        raise FileNotFoundError(f"Park adjustment workbook is missing: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    factors: dict[tuple[str, str], list[float | None]] = {}
    try:
        if season >= 2026:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "") for value in next(iterator)]
            for values in iterator:
                row = dict(zip(headers, values))
                code = str(row.get("Pitch") or "").strip()
                hb, ivb = _number(row.get("HB_Offset")), _number(row.get("IVB_Offset"))
                if code in PARK_FACTOR_CODES and hb is not None and ivb is not None:
                    factors[(_stadium(row.get("Stadium")), code)] = [hb, ivb]
        else:
            for metric, index in (("IVB", 1), ("HB", 0)):
                sheet = next(sheet for sheet in workbook.worksheets if metric in sheet.title.upper())
                for values in sheet.iter_rows(min_row=2, values_only=True):
                    stadium = _stadium(values[0])
                    for column, code in enumerate(PARK_FACTOR_CODES, 1):
                        value = _number(values[column] if column < len(values) else None)
                        if value is not None:
                            factors.setdefault((stadium, code), [None, None])[index] = value
    finally:
        workbook.close()
    return {key: (float(value[0]), float(value[1])) for key, value in factors.items() if None not in value}


def _movement_adjust(rows: list[dict], root: Path, season: int) -> dict:
    factors = _load_park_factors(root, season)
    available = adjusted = 0
    for row in rows:
        hb = _safe_float(row.get("horizontal_movement_cm"))
        ivb = _safe_float(row.get("vertical_movement_cm"))
        row["adjusted_hb_cm"], row["adjusted_ivb_cm"] = np.nan, np.nan
        if np.isnan(hb) or np.isnan(ivb):
            continue
        available += 1
        code = PARK_FACTOR_CODE.get(pitch_code(row), pitch_code(row))
        offset = factors.get((_stadium(row.get("stadium")), code))
        if offset is None:
            continue
        row["adjusted_hb_cm"] = hb + offset[0]
        row["adjusted_ivb_cm"] = ivb + offset[1]
        adjusted += 1
    return {
        "movement_available": available,
        "movement_adjusted": adjusted,
        "adjustment_coverage_pct": round(100 * adjusted / available, 4) if available else 0.0,
        "formula": "adjusted movement = Visual Baseball measurement + stadium/pitch offset",
    }


def predict_pzone(train: list[dict], test: list[dict], fields: tuple[str, ...] = PZONE_NUMERIC) -> np.ndarray:
    """Fit the take-only CalledStrike vs Ball/HBP model and score held-out pitches."""
    take = [row for row in train if row["decision_type"] == "Take"]
    target = np.array([str(row.get("pitch_call_code") or "").upper() == "T" or row.get("event") == "CalledStrike" for row in take], dtype=int)
    model = _classifier().fit(_encode_numeric(take, fields), target)
    probability = model.predict_proba(_encode_numeric(test, fields))[:, list(model.classes_).index(1)]
    return np.clip(probability, 1e-6, 1 - 1e-6)


def _write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(
            file, fieldnames=fields, extrasaction="ignore", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)
