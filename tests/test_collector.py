import json
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from visualbaseball.collector import process_payload
from visualbaseball.naver import NaverEnrichment, build_enrichment, pitch_key
from visualbaseball.parser import _half_key, parse_game
from visualbaseball.state_machine import GameState
from visualbaseball.curated import load_rows
from visualbaseball.storage import Store
from visualbaseball.validation import validate_game
from visualbaseball.export_excel import export_latest


ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "data" / "raw" / "2026" / "20260328HTSK0.json"


def test_sample_game_and_idempotency(tmp_path):
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    ok, message, pitches = process_payload(tmp_path, payload)
    assert ok and message == "PASS" and pitches == 338
    first = len(load_rows(tmp_path, "pitches", 2026))
    ok, message, pitches = process_payload(tmp_path, payload)
    assert ok and message == "PASS" and pitches == 338
    assert len(load_rows(tmp_path, "pitches", 2026)) == first
    rows = load_rows(tmp_path, "pitches", 2026)
    assert len({row["pitch_id"] for row in rows}) == 338
    assert not any("spin" in key.lower() for row in rows for key in row)
    game_stadium = load_rows(tmp_path, "games", 2026)[0]["stadium"]
    assert {row["stadium"] for row in rows} == {game_stadium}
    events = load_rows(tmp_path, "events", 2026)
    assert {row["stadium"] for row in events} == {game_stadium}
    assert not any(row["event_code"] == "OFFICIAL_LINESCORE_RECONCILIATION" for row in events)
    assert sum(row["runs_on_pitch"] for row in rows) == 13


def test_y0_catcher_and_naver_wp_pb_fields_are_retained():
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    first_half, first_pa = payload["pbpData"][0], payload["pbpData"][0]["pas"][0]
    key = pitch_key(first_half["inning"], first_half["half"], first_pa["batterId"], first_pa["pitcherId"], 1)
    enrichment = NaverEnrichment(
        payload["gameData"]["gameId"], "naver-game", ["https://example.test/relay"],
        {key: [{"naver_pitch_id": "pitch-1", "is_wild_pitch": True, "is_passed_ball": False}]},
        {"home": {"id": "catcher-home", "name": "홈포수", "source": "naver_lineup"}},
    )
    _, _, pitches, _ = parse_game(payload, naver_enrichment=enrichment)
    first = pitches[0]
    assert first["y0"] == first_pa["pitches"][0]["y0"]
    assert (first["catcher_id"], first["catcher_name"], first["catcher_source"]) == ("catcher-home", "홈포수", "naver_lineup")
    assert (first["is_wild_pitch"], first["is_passed_ball"], first["naver_pitch_id"], first["naver_match_status"]) == (True, False, "pitch-1", "matched")


def test_naver_relay_event_is_assigned_to_the_previous_pitch():
    payload = {"result": {"textRelayData": {
        "gameId": "20260328KTLG02026",
        "homeLineup": {"batter": [{"pos": 2, "seqno": 1, "pcode": "10", "name": "홈포수"}]},
        "awayLineup": {"batter": [{"pos": 2, "seqno": 1, "pcode": "20", "name": "원정포수"}]},
        "textRelays": [{"inn": 1, "homeOrAway": "0", "textOptions": [
            {"text": "1구 볼", "ptsPitchId": "p1", "pitchNum": 1, "currentGameState": {"batter": "1", "pitcher": "2"}},
            {"text": "1루주자 홍길동 : 폭투로 2루까지 진루", "ptsPitchId": None},
            {"text": "2구 볼", "ptsPitchId": "p2", "pitchNum": 2, "currentGameState": {"batter": "1", "pitcher": "2"}},
            {"text": "2루주자 홍길동 : 포일로 3루까지 진루", "ptsPitchId": None},
            {"text": "3구 볼", "ptsPitchId": "p3", "pitchNum": 3, "currentGameState": {"batter": "1", "pitcher": "2"}},
        ]}],
    }}}
    enrichment = build_enrichment("20260328KTLG0", [payload])
    assert enrichment.starters["home"]["id"] == "10"
    assert enrichment.pitch_events[pitch_key(1, "top", "1", "2", 1)][0]["is_wild_pitch"] is True
    assert enrichment.pitch_events[pitch_key(1, "top", "1", "2", 2)][0]["is_passed_ball"] is True
    assert enrichment.pitch_events[pitch_key(1, "top", "1", "2", 3)][0]["is_passed_ball"] is False


def test_completed_status_is_not_written_when_processed_tables_fail(tmp_path):
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    with patch.object(Store, "replace_game", side_effect=OSError("disk full")):
        with pytest.raises(OSError, match="disk full"):
            process_payload(tmp_path, payload)
    assert Store(tmp_path).manifest()["games"].get("20260328HTSK0") is None


def test_count_transitions_and_two_strike_foul():
    state = GameState(); state.apply_non_terminal_pitch("B"); assert state.balls == 1
    state.apply_non_terminal_pitch("S"); state.apply_non_terminal_pitch("F"); state.apply_non_terminal_pitch("F")
    assert state.strikes == 2


def test_runner_advances_multiple_outs_and_inning_reset():
    state = GameState(); state.set_bases({"b1": {"id": "a"}, "b2": {"id": "b"}}); state.outs = 1
    state.begin_half(2, "bottom")
    assert (state.outs, state.base_state_code, state.balls, state.strikes) == (0, 0, 0, 0)


def test_source_snapshot_at_half_start_is_retained_as_source_limited():
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    payload["pbpData"][0]["pas"][0]["basesBefore"] = {"b1": {"id": "runner"}, "b2": None, "b3": None}
    game, events, pitches, _ = parse_game(payload)
    assert pitches[0]["parse_status"] == "source_limited"
    assert validate_game(game, events, pitches) == (True, "PASS")
    assert any(event["event_code"] == "SOURCE_SNAPSHOT" and event["parse_status"] == "unknown" for event in events)


def test_source_halves_are_ordered_chronologically_before_parsing():
    assert sorted([{"inning": 7, "half": "bottom"}, {"inning": 6, "half": "bottom"}, {"inning": 7, "half": "top"}], key=_half_key) == [
        {"inning": 6, "half": "bottom"}, {"inning": 7, "half": "top"}, {"inning": 7, "half": "bottom"},
    ]


def test_incremental_skip_logic(tmp_path):
    store = Store(tmp_path); raw = tmp_path / "data/raw/2026/old.json"; raw.parent.mkdir(parents=True); raw.write_text("{}")
    store.mark("old", "completed", raw, "PASS")
    assert store.should_fetch("old", "2000-01-01") is False


def test_incremental_skip_recovers_portable_raw_cache_from_legacy_absolute_path(tmp_path):
    store = Store(tmp_path)
    raw = tmp_path / "data/raw/2026/game.json"; raw.parent.mkdir(parents=True); raw.write_text("{}")
    store.mark("game", "completed", Path("C:/another-machine/data/raw/2026/game.json"), "PASS")
    assert store.should_fetch("game", "2026-07-26") is False


def test_official_linescore_overrides_conflicting_pbp_snapshot():
    payload = json.loads((ROOT / "data/raw/2026/20260527HTWO0.json").read_text(encoding="utf-8"))
    game, events, pitches, _ = parse_game(payload)
    assert validate_game(game, events, pitches) == (True, "PASS")
    assert (events[-1]["away_score_after"], events[-1]["home_score_after"]) == (9, 2)
    conflicts = [event for event in events if event["event_code"] == "SOURCE_SCORE_CONFLICT"]
    assert conflicts and all(event["parse_status"] == "unknown" for event in conflicts)
    assert not any(event["event_code"] == "OFFICIAL_LINESCORE_RECONCILIATION" for event in events)


def test_season_export_uses_separate_storage_and_output_name(tmp_path):
    storage_root, output_root = tmp_path / "season", tmp_path / "output"
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    assert process_payload(storage_root, payload, season=2025, curated_root=output_root)[0]
    output = export_latest(output_root, 2025)
    assert output.name == "visualbaseball_savant_2025_latest.xlsx"
    assert output.exists()
    workbook = load_workbook(output, read_only=True)
    assert workbook["Pitches"].max_row == 339
    workbook.close()


def _first_nonterminal_ball(payload):
    pa_counter = 0
    for half in payload["pbpData"]:
        for pa in half.get("pas") or []:
            pa_counter += 1
            pitches = pa.get("pitches") or []
            for index, pitch in enumerate(pitches[:-1], 1):
                if pitch.get("r") == "B":
                    game_id = payload["gameData"]["gameId"]
                    return f"{game_id}-{game_id}-{pa_counter:03d}-{index:02d}", pa, pitch


def test_trackman_bunt_foul_correction_changes_call_and_count_only_when_row_matches(monkeypatch):
    # Every row carries a wall-clock fetched_at; pin it so parses that straddle a second compare equal.
    monkeypatch.setattr("visualbaseball.parser._now", lambda: "2026-01-01T00:00:00+00:00")
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    pitch_id, pa, pitch = _first_nonterminal_ball(payload)
    fix = {"pitch_id": pitch_id, "batter_id": str(pa["batterId"]), "pitcher_id": str(pa["pitcherId"]),
           "source_code": "B", "code": "W", "source_velocity_kmh": float(pitch["spd"])}
    base = {row["pitch_id"]: row for row in parse_game(payload, season=2026)[2]}
    fixed_game, fixed_events, fixed_pitches, _ = parse_game(payload, season=2026, call_corrections={pitch_id: fix})
    fixed = {row["pitch_id"]: row for row in fixed_pitches}
    row, nxt = fixed[pitch_id], fixed[pitch_id[:-2] + f"{int(pitch_id[-2:]) + 1:02d}"]
    assert row["pitch_call_code"] == "W" and row["description"] == "Bunt Foul"
    assert not row["is_swing"] and not row["is_take"] and not row["is_contact"]
    assert (row["balls_after"], row["strikes_after"]) == (row["balls_before"], row["strikes_before"] + 1)
    assert (nxt["balls_before"], nxt["strikes_before"]) == (row["balls_before"], row["strikes_before"] + 1)
    assert validate_game(fixed_game, fixed_events, fixed_pitches)[0]
    changed = {pid for pid in base if base[pid] != fixed[pid]}
    assert changed and all(pid.rsplit("-", 1)[0] == pitch_id.rsplit("-", 1)[0] for pid in changed)
    stale = {**fix, "source_velocity_kmh": fix["source_velocity_kmh"] + 1}
    assert parse_game(payload, season=2026, call_corrections={pitch_id: stale})[2] == parse_game(payload, season=2026)[2]


def test_count_correction_sets_start_and_inserts_calls_only_when_the_pa_matches(monkeypatch):
    monkeypatch.setattr("visualbaseball.parser._now", lambda: "2026-01-01T00:00:00+00:00")
    payload = json.loads(FIXTURE.read_text(encoding="utf-8-sig"))
    pitch_id, pa, _ = _first_nonterminal_ball(payload)
    pa_id = pitch_id.rsplit("-", 1)[0].split("-", 1)[1]
    codes = "".join(str(p.get("r", "")).upper() for p in pa["pitches"])
    fix = {"pa_id": pa_id, "batter_id": str(pa["batterId"]), "pitcher_id": str(pa["pitcherId"]), "source_codes": codes,
           "start": [0, 1], "inserts": [{"before_pitch": 2, "code": "B"}]}
    base = parse_game(payload, season=2026)[2]
    fixed = parse_game(payload, season=2026, count_corrections={pa_id: fix})[2]
    rows = [r for r in fixed if r["pa_id"] == pa_id]
    assert (rows[0]["balls_before"], rows[0]["strikes_before"]) == (0, 1)
    changed = {b["pitch_id"] for b, f in zip(base, fixed) if b != f}
    assert changed and all(pid.rsplit("-", 1)[0].endswith(pa_id) for pid in changed)
    stale = {**fix, "source_codes": codes + "B"}
    assert parse_game(payload, season=2026, count_corrections={pa_id: stale})[2] == base
